from pptx import Presentation
import sys

prs = Presentation('美股流动性与基本面双轮驱动分析.pptx')
print(f'Total slides: {len(prs.slides)}')
print('='*60)

for i, slide in enumerate(prs.slides):
    print(f'\n--- Slide {i+1} ---')
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text.strip():
            text = shape.text.strip()[:500]
            print(text)
